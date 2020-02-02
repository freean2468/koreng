import scrapy
import re
import json
import io
import StringIO
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl import load_workbook
import datetime

xlFile = "sample.xlsx"

class ScriptSpider(scrapy.Spider):
    name = "scripts"
    #'https://www.springfieldspringfield.co.uk/'
    # 'https://www.gurufocus.com/stock_list.php?m_country[]=USA&p=0&n=30'
    #"http://us.itooza.com/stocks/summary/TSLA"
    start_urls = [
        
    ]

    data = []
    wb = Workbook()
    sheet = wb.worksheets[0]

    def __init__(self):
        self.path = '/Users/hoon-ilsong/Documents/crawler products/'

        # Rows can also be appended
        self.sheet.append([datetime.datetime.now(), "re", 10, "", "", 'Earnings Per Share', "", 'Price Earning Ratio', "", "", "Book-value per Share", "Price to Book Ratio", "", "Return on Equity", "", "", "", "", "Residual Income Model"])
        self.sheet.append(["code",	"price",	"shares outstanding",	"net income(A)",	"Y",	"net income(Q)",	"EPS(A)",	"EPS(Q)",	"PER(A)",	"PER(Q)",	"Total Assets",	"Total liabilities",	"BPS(Q)",	"PBR(Q)",	"Total Equity",	"ROE(Q)",	"ROE(A)",	"P/P",	"P", 	"5.00", 	"0.00", 	"-5"])

        with open('stock_list.txt', 'r') as f:
            stock_list = f.read()
            stock_list = stock_list.splitlines()
            
            idx = 0
            for code in stock_list:
                idx = idx+1
                self.start_urls.append("http://us.itooza.com/stocks/summary/"+code)
                # if(idx > 200):
                #     break

        # crawling stock list
        # idx = 0
        # while idx < 686:
        #     idx = idx + 1
        #     self.start_urls.append('https://www.gurufocus.com/stock_list.php?m_country[]=USA&p=%d&n=30' % idx)

    def parse(self, response):
        try:
            if "/summary/" in response.url: 

                outstanding = response.xpath("//td/text()")[2].get()
                # print(outstanding)

                Y = response.xpath("//span/text()")[59].get()
                # print(Y)

                code = response.xpath("//span/text()")[5].get()
                # print(code)

                price = response.xpath("//li/text()")[7].get()
                # print(price)

                ANI = response.xpath("//td/text()")[34].get() + "000000"
                # print(ANI)

                QNI = response.xpath("//td/text()")[37].get() + "000000"
                # print(QNI)

                # for node in response.xpath("//span/text()"):
                #     print(node)

                # for node in response.xpath("//td/text()"):
                #     print(node)

                # for idx, row in enumerate(self.sheet.rows):
                #     # print(row[0].value)
                #     if row[0].value == code:
                #         # print("I'm in!")
                        
                #         row[1].value = price
                #         row[2].value = outstanding
                #         row[3].value = ANI
                #         row[4].value = Y
                #         row[5].value = QNI
                #         return

                # Rows can also be appended
                self.sheet.append([code, price,	outstanding, ANI, Y, QNI, "=D3/$C$3",	"=F3/$C$3",	"=$B$3/G3",	"=$B$3/H3",	"Total Assets",	"Total liabilities",	"=(K3-L3)/C3",	"=B3/M3",	"Total Equity",	"=F3/O3*100",	"=D3/O3*100",	"=B3/S3",	"=($O3+$O3*($Q3-$C$1)/$C$1)/$C3", 	"=($O3+$O3*($Q3+T$2-$C$1)/$C$1)/$C3", 	"=($O3+$O3*($Q3+U$2-$C$1)/$C$1)/$C3", 	"=($O3+$O3*($Q3+V$2-$C$1)/$C$1)/$C3"])

                next_page = response.urljoin(response.url.replace('/summary/', '/financials/'))
                yield scrapy.Request(next_page, callback=self.parse)
                

            else:
                code = response.xpath("//span/text()")[5].get()
                # print('code : ' + code)

                # for node in response.xpath("//td/text()"):
                #     print(node)
                
                asset = ""
                liability = ""
                equity = ""

                for row in response.xpath('//table[@class="tableColtype typeScroll tableReports"]/tbody/tr'):
                    asset = row.xpath('//tr[@class="totalAsset"]/td[2]//text()').extract_first() + "000000"
                    liability = row.xpath('//tr[@class="totalAsset"][2]/td[2]//text()').extract_first() + "000000"
                    equity = row.xpath('//tr[@class="totalAsset"][3]/td[2]//text()').extract_first() + "000000"
                    break

                # for node in response.xpath("//span/text()"):
                #     print(node)

                # for node in response.xpath("//td/text()"):
                #     print(node)

                for idx, row in enumerate(self.sheet.rows):
                    # print(row[0].value)
                    if row[0].value == code:
                        # print("I'm in!")
                        
                        row[10].value = asset
                        row[11].value = liability
                        row[14].value = equity
                        self.wb.save(xlFile)
                        print("code : "+code+", saved!")
                        return

                # self.sheet.append([code, "price",	"outstanding", "ANI", "Y", "QNI", "=D3/$C$3", "=F3/$C$3",	"=$B$3/G3",	"=$B$3/H3",	asset, liability,	"=(K3-L3)/C3",	"=B3/M3",	equity,	"=F3/O3*100",	"=D3/O3*100",	"=B3/S3",	"=($O3+$O3*($Q3-$C$1)/$C$1)/$C3", 	"=($O3+$O3*($Q3+T$2-$C$1)/$C$1)/$C3", 	"=($O3+$O3*($Q3+U$2-$C$1)/$C$1)/$C3", 	"=($O3+$O3*($Q3+V$2-$C$1)/$C$1)/$C3"])
        except Exception as ex:
            # print('ex : ' + ex + ', No code : ' + response.url)
            print('No code : ' + response.url)



                # data.append([row[0].value, row[21].value])

            # print(data)

        # ''' stock list crawling'''
        # data = []
        # for row in response.xpath('//*[@class="R5"]//tbody/tr'):
        #     code = row.xpath('td[1]//text()').extract_first()
        #     data.append(code)

        # with open('stock_list.txt', 'r') as f:
        #     stock_list = f.read()
        #     stock_list = stock_list.splitlines()
        #     stock_list = stock_list + data
            
        #     for code in stock_list:
        #         print(code)

        #     with open('stock_list.txt', 'w') as f:
        #         for item in stock_list:
        #             f.write("%s\n" % item)

'''Drama Script
https://www.springfieldspringfield.co.uk/view_episode_scripts.php?tv-show=spartacus-gods-of-the-arena&episode=s01e01
'''
        # page = response.url.split("/")[-2]
        # document = Document()

        # # u'The Vampire Diaries s04e15 Episode Script            '
        # h = response.css("div.main-content-left h1::text").get()
        # title = re.sub("[']", '', h).strip()
        # # title = title.replace('u', '', 1)

        # subtitle = response.css("div.main-content-left h3::text").get()
        # document.add_heading(subtitle, 0)

        # res = [i for i in range(len(title)) if title.startswith(' ', i)]
        # title = title[0:res[-2]]

        # body = response.css("div.scrolling-script-container::text").getall()

        # document.add_paragraph(body)
        # document.save(self.path+title+'.docx')

        # self.log('<------------------------------ Saved file %s ----------------------------->' % title)

        # next_page = ''

        # i = 0
        # PorN = response.css("div.episode_script a::text").getall()
        # while(i < len(PorN)):
        #     _temp = PorN[i]
        #     _temp = re.sub("[']", '', _temp)
        #     self.log(_temp)
        #     if _temp == 'Next Episode':
        #         next_page = response.css("div.episode_script a::attr(href)").getall()[i]
        #         break
        #     i=i+1

        # if next_page is not None:
        #     next_page = response.urljoin(next_page)
        #     yield scrapy.Request(next_page, callback=self.parse)
